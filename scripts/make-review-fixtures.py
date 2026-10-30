# 生成审阅稿写回测试用的合成 docx（全部虚构内容）。
# 用法：python3 scripts/make-review-fixtures.py [输出目录，默认 work/review-fixtures]
# 每个用例一对：<名字>-A.docx（我方原稿）/ <名字>-B.docx（对方改稿）。
import sys, os, copy, random
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from lxml import etree

OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(__file__), '..', 'work', 'review-fixtures')
os.makedirs(OUT, exist_ok=True)
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
NS = f'xmlns:w="{W}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'


def esc(s):
    return s.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')


def r(t, b=False, i=False, sz=None):
    pr = ''
    if b: pr += '<w:b/>'
    if i: pr += '<w:i/>'
    if sz: pr += f'<w:sz w:val="{sz * 2}"/>'
    return f'<w:r>{"<w:rPr>" + pr + "</w:rPr>" if pr else ""}<w:t xml:space="preserve">{esc(t)}</w:t></w:r>'


def p(*runs, style=None, num=None):
    ppr = ''
    if style: ppr += f'<w:pStyle w:val="{style}"/>'
    if num is not None: ppr += f'<w:numPr><w:ilvl w:val="0"/><w:numId w:val="{num}"/></w:numPr>'
    body = ''.join(x if x.startswith('<') else r(x) for x in runs)
    return f'<w:p>{"<w:pPr>" + ppr + "</w:pPr>" if ppr else ""}{body}</w:p>'


def tbl(rows):
    cols = len(rows[0])
    grid = ''.join('<w:gridCol w:w="3000"/>' for _ in range(cols))
    trs = ''
    for row in rows:
        tcs = ''.join(f'<w:tc><w:tcPr><w:tcW w:w="3000" w:type="dxa"/></w:tcPr>{p(c)}</w:tc>' for c in row)
        trs += f'<w:tr>{tcs}</w:tr>'
    return f'<w:tbl><w:tblPr><w:tblStyle w:val="TableGrid"/><w:tblW w:w="0" w:type="auto"/></w:tblPr><w:tblGrid>{grid}</w:tblGrid>{trs}</w:tbl>'


def make(name, side, body_xml, header=None, comments=None, numbering=False):
    d = Document()
    if numbering:
        # 让模板带上 numbering 部件
        d.add_paragraph('x', style='List Number')
    if header is not None:
        sec = d.sections[0]
        sec.header.paragraphs[0].text = header
    body = d.element.body
    sect = body.find(qn('w:sectPr'))
    for c in list(body):
        if c is not sect:
            body.remove(c)
    frag = etree.fromstring(f'<w:body {NS}>{body_xml}</w:body>')
    for c in list(frag):
        sect.addprevious(c)
    path = os.path.join(OUT, f'{name}-{side}.docx')
    d.core_properties.author = '虚构作者'
    d.save(path)
    if comments:
        add_comments(path, comments)
    return path


def add_comments(path, comments):
    # 给已有的批注（对方文件里原来就有的）：comments = [(id, author, text)]
    import zipfile, shutil
    tmp = path + '.tmp'
    zin = zipfile.ZipFile(path)
    zout = zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED)
    for it in zin.infolist():
        data = zin.read(it.filename)
        if it.filename == '[Content_Types].xml':
            data = data.replace(b'</Types>', b'<Override PartName="/word/comments.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"/></Types>')
        if it.filename == 'word/_rels/document.xml.rels':
            data = data.replace(b'</Relationships>', b'<Relationship Id="rIdCm1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments" Target="comments.xml"/></Relationships>')
        zout.writestr(it, data)
    cm = ''.join(f'<w:comment w:id="{i}" w:author="{a}" w:date="2026-10-01T00:00:00Z" w:initials="X"><w:p><w:r><w:t>{esc(t)}</w:t></w:r></w:p></w:comment>' for i, a, t in comments)
    zout.writestr('word/comments.xml', f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:comments xmlns:w="{W}">{cm}</w:comments>')
    zout.close(); zin.close()
    shutil.move(tmp, path)


# ---------- 1. 基础：改词、增删、整句替换、混合样式 ----------
A = [
    p('第一条 定义', style='Heading1'),
    p('本协议中“服务”指乙方向甲方提供的', r('云存储与备份', b=True), '服务。'),
    p('服务期限为十二个月，自本协议生效之日起计算。'),
    p('甲方应在收到发票后三十日内支付服务费。'),
    p('任何一方不得向第三方披露本协议内容。'),
    p('本条款适用于双方的关联公司。'),
    p('第二条 费用', style='Heading1'),
    p('服务费总额为人民币 120,000 元，分两期支付。'),
]
B = [
    p('第一条 定义', style='Heading1'),
    p('本协议中“服务”指乙方向甲方提供的', r('云存储、备份与恢复', b=True), '服务。'),
    p('服务期限为二十四个月，自本协议生效之日起计算。'),
    p('甲方应在收到发票后六十日内支付全部服务费。'),
    p('本条款适用于双方的关联公司。'),
    p('乙方可在提前三十日书面通知后调整服务内容。'),
    p('第二条 费用', style='Heading1'),
    p('服务费总额为人民币 198,000 元，一次性支付。'),
]
make('basic', 'A', ''.join(A)); make('basic', 'B', ''.join(B))

# ---------- 2. 标题与列表 ----------
A = [p('第三条 保密', style='Heading1'), p('保密信息包括：', style=None),
     p('技术资料', style='ListNumber'), p('商业计划', style='ListNumber'), p('客户名单', style='ListNumber'),
     p('第四条 违约', style='Heading1'), p('违约方应赔偿守约方的全部损失。')]
B = [p('第三条 保密义务', style='Heading1'), p('保密信息包括：'),
     p('技术资料', style='ListNumber'), p('客户名单及联系方式', style='ListNumber'),
     p('第四条 违约', style='Heading2'), p('违约方应赔偿守约方的直接损失，但最高不超过合同总额的 10%。'), p('第五条 其他', style='Heading1')]
make('lists', 'A', ''.join(A), numbering=True); make('lists', 'B', ''.join(B), numbering=True)

# ---------- 3. 表格：改单元格、加行、删行 ----------
A = [p('报价明细如下：'), tbl([['项目', '数量', '单价'], ['存储', '10 TB', '800'], ['备份', '5 TB', '600'], ['恢复', '按次', '300']]), p('以上价格含税。')]
B = [p('报价明细如下：'), tbl([['项目', '数量', '单价'], ['存储', '20 TB', '750'], ['恢复', '按次', '300'], ['迁移', '1 次', '5,000']]), p('以上价格含税，不含差旅。')]
make('table', 'A', ''.join(A)); make('table', 'B', ''.join(B))

# ---------- 4. 超链接、域、制表符、换行 ----------
link = lambda t: f'<w:hyperlink w:anchor="svc" w:history="1">{r(t)}</w:hyperlink>'
fld = lambda res: f'<w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve"> DATE </w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r>{r(res)}<w:r><w:fldChar w:fldCharType="end"/></w:r>'
tabbed = lambda a, b: f'<w:r><w:t>{a}</w:t><w:tab/><w:t>{b}</w:t><w:br/><w:t>附注</w:t></w:r>'
A = [p('详见附件一及', link('服务说明网页'), '。'), p('签署日期：', fld('2026年3月1日')), p(tabbed('甲方', '乙方')), p('一行之内有', r('多个'), r('样式', i=True), '的', r('文字', b=True), '。')]
B = [p('详见附件二及', link('服务说明与价格网页'), '。'), p('签署日期：', fld('2026年9月1日')), p(tabbed('甲方', '乙方（盖章）')), p('一行之内有', r('若干'), r('样式', i=True), '的', r('文字段落', b=True), '。')]
make('inline', 'A', ''.join(A)); make('inline', 'B', ''.join(B))

# ---------- 5. 对方文件里原本就有修订和批注 ----------
A = [p('付款方式为银行转账。'), p('发票类型为增值税普通发票。'), p('本合同一式两份。')]
ins = lambda t: f'<w:ins w:id="801" w:author="对方律师" w:date="2026-10-01T00:00:00Z">{r(t)}</w:ins>'
dele = lambda t: f'<w:del w:id="802" w:author="对方律师" w:date="2026-10-01T00:00:00Z"><w:r><w:delText xml:space="preserve">{esc(t)}</w:delText></w:r></w:del>'
cmt = lambda i, t: f'<w:commentRangeStart w:id="{i}"/>{r(t)}<w:commentRangeEnd w:id="{i}"/><w:r><w:commentReference w:id="{i}"/></w:r>'
B = [p('付款方式为', dele('银行转账'), ins('银行转账或承兑汇票'), '。'), p('发票类型为', cmt(0, '增值税专用发票'), '。'), p('本合同一式四份。')]
make('prior', 'A', ''.join(A)); make('prior', 'B', ''.join(B), comments=[(0, '对方律师', '请确认税率')])

# ---------- 6. 页眉变化 + 分节 ----------
A = [p('正文第一段保持不变。'), p('正文第二段将被修改。')]
B = [p('正文第一段保持不变。'), p('正文第二段已经被修改。')]
make('header', 'A', ''.join(A), header='虚构科技 · 草稿'); make('header', 'B', ''.join(B), header='虚构科技 · 第二稿')

# ---------- 7. 文本框里的字被改（应判“没写进”） ----------
def txbx(t):
    return ('<w:r><mc:AlternateContent><mc:Choice Requires="wps"><w:drawing><wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="1" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">'
            '<wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="column"><wp:posOffset>0</wp:posOffset></wp:positionH><wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset></wp:positionV>'
            '<wp:extent cx="1800000" cy="400000"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/><wp:docPr id="1" name="TB1"/><wp:cNvGraphicFramePr/>'
            '<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:wsp><wps:cNvSpPr txBox="1"/><wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="1800000" cy="400000"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></wps:spPr>'
            f'<wps:txbx><w:txbxContent>{p(t)}</w:txbxContent></wps:txbx><wps:bodyPr/></wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice></mc:AlternateContent></w:r>')
A = [p('封面说明：', txbx('机密 · 仅限内部')), p('普通段落一。')]
B = [p('封面说明：', txbx('机密 · 双方共享')), p('普通段落一（已更新）。')]
make('textbox', 'A', ''.join(A)); make('textbox', 'B', ''.join(B))

# ---------- 8. 大文档（性能）：约 100 页 ----------
random.seed(7)
words = '甲方 乙方 服务 期限 费用 保密 违约 责任 通知 争议 解决 适用 法律 附件 生效 终止 数据 安全 交付 验收 支付 发票 税费 知识产权 许可'.split()
def sent(k):
    return '，'.join(''.join(random.choice(words) for _ in range(random.randint(3, 6))) for _ in range(k)) + '。'
A, B = [], []
for i in range(1, 401):
    if i % 20 == 1:
        h = p(f'第{i // 20 + 1}章 条款', style='Heading1'); A.append(h); B.append(h)
    a = sent(random.randint(3, 6))
    if i % 7 == 0:
        bb = a.replace('，', '，并且', 1)
    elif i % 11 == 0:
        bb = sent(4)
    else:
        bb = a
    if i % 37 != 0: A.append(p(a))
    if i % 41 != 0: B.append(p(bb))
make('large', 'A', ''.join(A)); make('large', 'B', ''.join(B))
# 约 100 页：同样的规律放大 4 倍，并夹进表格
A2, B2 = [], []
for k in range(4):
    A2 += A; B2 += B
    t = [['项目', '数量', '单价']] + [[f'条目{k}-{j}', f'{j} 件', f'{100 + j}'] for j in range(12)]
    t2 = [row[:] for row in t]; t2[3][2] = '999'; t2[7][0] = '条目改名'
    A2.append(tbl(t)); B2.append(tbl(t2))
make('huge', 'A', ''.join(A2)); make('huge', 'B', ''.join(B2))

# ---------- 9. 结构：书签、内容控件、分节段落、移动、拆分合并 ----------
bm = lambda t: f'<w:bookmarkStart w:id="90" w:name="term"/>{r(t)}<w:bookmarkEnd w:id="90"/>'
sdt = lambda t: f'<w:sdt><w:sdtPr><w:alias w:val="金额"/></w:sdtPr><w:sdtContent>{r(t)}</w:sdtContent></w:sdt>'
sect = '<w:p><w:pPr><w:sectPr><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="1440" w:header="851" w:footer="992" w:gutter="0"/></w:sectPr></w:pPr>' + r('第一节结束。') + '</w:p>'
sect2 = '<w:p><w:pPr><w:sectPr><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="1440" w:header="851" w:footer="992" w:gutter="0"/></w:sectPr></w:pPr>' + r('第一节到此结束。') + '</w:p>'
A = [p('定义见', bm('服务期限'), '条款。'), p('合同金额：', sdt('壹拾万元'), '。'), sect,
     p('甲段：交付物包括源代码和文档。'), p('乙段：验收在交付后十个工作日内完成。'), p('丙段：本段将被挪到最后，内容保持完全一致。'),
     p('丁段：争议提交仲裁。这一句和下一句原来在同一段。下一句在这里。')]
B = [p('定义见', bm('服务期限及续约'), '条款。'), p('合同金额：', sdt('壹拾伍万元'), '。'), sect2,
     p('甲段：交付物包括可执行程序和文档。'), p('乙段：验收在交付后十个工作日内完成。'),
     p('丁段：争议提交法院诉讼。这一句和下一句原来在同一段。'), p('下一句在这里。'), p('丙段：本段将被挪到最后，内容保持完全一致。')]
make('struct', 'A', ''.join(A)); make('struct', 'B', ''.join(B))

# ---------- 10. 空文档边界：B 全新、A 全删 ----------
make('empty', 'A', p('只有一段旧文字。')); make('empty', 'B', p('全部换成了新的文字，没有一个字相同。') + p('还多了一段。'))

# ---------- 11. 模板识别：模板两个版本 + 对方在旧版上改的稿 ----------
base = [p(f'第{i}条 ' + t) for i, t in enumerate([
    '本合同由甲乙双方在平等自愿的基础上订立。', '乙方向甲方提供云存储服务，服务内容见附件一。', '服务期限为十二个月，自合同生效之日起计算。',
    '服务费按季度结算，甲方在收到发票后三十日内付款。', '乙方应保证数据的完整性与可用性，年可用率不低于百分之九十九点九。',
    '任何一方违约的，应向守约方支付违约金，违约金为合同总额的百分之二十。', '双方对在履行本合同过程中知悉的对方信息负有保密义务。',
    '因本合同引起的争议，双方应友好协商；协商不成的，提交甲方所在地人民法院诉讼解决。', '本合同一式两份，双方各执一份，具有同等法律效力。',
    '未尽事宜，双方可另行签订补充协议。', '附件一：服务说明。', '附件二：服务水平协议。'], 1)]
v2 = list(base)
v2[2] = p('第3条 服务期限为二十四个月，自合同生效之日起计算，期满自动续约一年。')
v2[5] = p('第6条 任何一方违约的，应向守约方支付违约金，违约金为合同总额的百分之十五。')
v2[9] = p('第10条 未尽事宜，双方可另行签订补充协议，补充协议与本合同具有同等效力。')
v2[1] = p('第2条 乙方向甲方提供云存储与备份服务，服务内容见附件一及附件三。')
vendor = list(base)
vendor[3] = p('第4条 服务费按年度一次性预付，甲方在合同签订后十日内付款。')
vendor[7] = p('第8条 因本合同引起的争议，提交乙方所在地仲裁委员会仲裁。')
make('tpl', 'v1', ''.join(base)); make('tpl', 'v2', ''.join(v2)); make('tpl', 'vendor', ''.join(vendor))
import time
for n, ts in [('v1', '2026-03-02'), ('v2', '2026-09-05'), ('vendor', '2026-10-20')]:
    tt = time.mktime(time.strptime(ts, '%Y-%m-%d')); os.utime(os.path.join(OUT, f'tpl-{n}.docx'), (tt, tt))

# ---------- 12. 往来稿：四轮来回 ----------
def neg(fee, pay, court, extra=None):
    x = [p('第一条 服务内容', style='Heading1'), p('乙方向甲方提供云存储服务。'),
         p('第五条 付款', style='Heading1'), p(f'甲方应在收到发票后{pay}日内付款。'),
         p('第八条 违约责任', style='Heading1'), p(f'违约金为合同总额的{fee}。'),
         p('第十二条 争议解决', style='Heading1'), p(f'争议提交{court}解决。')]
    if extra: x.append(p(extra))
    return ''.join(x)
make('neg', 'V1', neg('百分之二十', '三十', '甲方所在地人民法院'))
make('neg', 'V2', neg('百分之十五', '六十', '甲方所在地人民法院'))
make('neg', 'V3', neg('百分之二十', '六十', '甲方所在地人民法院'))
make('neg', 'V4', neg('百分之十五', '六十', '乙方所在地仲裁委员会', '本合同自双方签字盖章之日起生效。'))
print('ok', OUT)
