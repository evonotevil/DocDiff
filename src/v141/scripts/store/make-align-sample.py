from docx import Document
from docx.shared import Pt

P1 = ('The following terms of service (the "Terms") will apply, in addition to those set forth in the '
      'Northwind Partner Distribution Agreement currently available at '
      'https://northwind.example/partner-agreement, as it may be updated from time to time by '
      'Northwind or available at any successor URLs (the "DDA") and the GPG Addendum, to Developer\'s participation in the '
      'Participating Games Earnback Program ("Earnback"). If there is any conflict between these Terms and the DDA or '
      'GPG Addendum with respect to the Earnback, these Terms will prevail.')
P2 = ('Developer must submit a complete application no later than thirty (30) days after the Program start date. '
      'Northwind may, in its sole discretion, approve or reject any application.')
P3 = ('Earnback amounts will be calculated monthly and paid within sixty (60) days after the end of each calendar month, '
      'subject to the eligibility requirements described in Section 4.')


def base(d, title):
    d.add_paragraph('NORTHWIND STUDIO · CONFIDENTIAL')
    d.add_paragraph('Contract ID: n244732')
    d.add_heading(title, 1)


# A：一段完整的段落（将转成 PDF）
a = Document()
base(a, 'PARTNER GAMES EARNBACK TERMS OF SERVICE')
a.add_paragraph(P1)
a.add_paragraph(P2)
a.add_paragraph(P3)
a.save('/home/claude/samples_store/split_A.docx')

# B：同样的内容，但第一段被拆成两段；第三段改了一个数字
b = Document()
base(b, 'PARTNER GAMES EARNBACK TERMS OF SERVICE')
cut = P1.index('Northwind Partner')
b.add_paragraph(P1[:cut].rstrip())
b.add_paragraph(P1[cut:])
b.add_paragraph(P2)
b.add_paragraph(P3.replace('sixty (60)', 'forty-five (45)'))
b.save('/home/claude/samples_store/split_B.docx')
print('ok')
